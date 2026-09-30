"""Deterministic Synthetic Dataset Generator for AI Resume Analyzer Evaluation.

Generates:
  - 100 synthetic resumes (TXT for all, plus PDF and DOCX for multi-format validation)
  - 20 synthetic job descriptions (high, medium, low, and zero-skill density)
  - 320 resume-JD evaluation pairs
  - Controlled ground-truth gold annotations (derived strictly BEFORE production parsing)

DISCLAIMER:
  ALL DATA GENERATED IS SYNTHETIC — NOT REAL RESEARCH DATA.
  NO REAL PEOPLE'S RESUMES OR PERSONAL INFORMATION ARE USED.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import docx
import pymupdf

# Skill Category Mapping from backend/app/database/skills.json
SKILL_TO_CATEGORY: dict[str, str] = {
    # Programming Languages
    "Python": "Programming Languages", "Java": "Programming Languages",
    "JavaScript": "Programming Languages", "TypeScript": "Programming Languages",
    "C": "Programming Languages", "C++": "Programming Languages", "C#": "Programming Languages",
    "Go": "Programming Languages", "Rust": "Programming Languages", "PHP": "Programming Languages",
    "R": "Programming Languages", "Kotlin": "Programming Languages", "Swift": "Programming Languages",
    "Dart": "Programming Languages", "DSA": "Programming Languages", "Data Structures": "Programming Languages",
    # Frontend
    "HTML": "Frontend", "HTML5": "Frontend", "CSS": "Frontend", "CSS3": "Frontend",
    "Bootstrap": "Frontend", "Tailwind CSS": "Frontend", "Tailwind": "Frontend",
    "React": "Frontend", "Next.js": "Frontend", "Angular": "Frontend", "Vue.js": "Frontend",
    "Redux": "Frontend", "jQuery": "Frontend",
    # Backend
    "Node.js": "Backend", "Node": "Backend", "Express.js": "Backend", "Spring Boot": "Backend",
    "FastAPI": "Backend", "Flask": "Backend", "Django": "Backend", ".NET": "Backend", "Laravel": "Backend",
    # Database
    "SQL": "Database", "MySQL": "Database", "PostgreSQL": "Database", "MongoDB": "Database",
    "SQLite": "Database", "Oracle": "Database", "Redis": "Database", "Firebase": "Database", "Supabase": "Database",
    # AI & ML
    "AI": "AI & ML", "Artificial Intelligence": "AI & ML", "Machine Learning": "AI & ML",
    "Deep Learning": "AI & ML", "Generative AI": "AI & ML", "GenAI": "AI & ML",
    "TensorFlow": "AI & ML", "PyTorch": "AI & ML", "Scikit-learn": "AI & ML", "OpenCV": "AI & ML",
    "NLP": "AI & ML", "spaCy": "AI & ML", "Transformers": "AI & ML", "XGBoost": "AI & ML",
    "Pandas": "AI & ML", "NumPy": "AI & ML", "Matplotlib": "AI & ML", "Seaborn": "AI & ML",
    "LangChain": "AI & ML", "RAG": "AI & ML", "Prompt Engineering": "AI & ML",
    "Exploratory Data Analysis": "AI & ML", "EDA": "AI & ML", "Regression": "AI & ML",
    "Classification": "AI & ML", "Clustering": "AI & ML",
    # Cloud & DevOps
    "AWS": "Cloud & DevOps", "Azure": "Cloud & DevOps", "Google Cloud": "Cloud & DevOps",
    "GCP": "Cloud & DevOps", "Docker": "Cloud & DevOps", "Kubernetes": "Cloud & DevOps",
    "Jenkins": "Cloud & DevOps", "Terraform": "Cloud & DevOps", "CI/CD": "Cloud & DevOps",
    # BI & Analytics
    "PowerBI": "BI & Analytics", "Power BI": "BI & Analytics", "Tableau": "BI & Analytics",
    "Excel": "BI & Analytics", "Looker": "BI & Analytics", "Google Analytics": "BI & Analytics",
    # Tools & IDEs
    "Git": "Tools & IDEs", "GitHub": "Tools & IDEs", "GitLab": "Tools & IDEs",
    "Bitbucket": "Tools & IDEs", "Jupyter Notebook": "Tools & IDEs", "Jupyter": "Tools & IDEs", "VS Code": "Tools & IDEs",
    # Testing & Quality
    "Testing": "Testing & Quality", "Software Testing": "Testing & Quality",
    "JUnit": "Testing & Quality", "PyTest": "Testing & Quality", "Selenium": "Testing & Quality",
    "Postman": "Testing & Quality", "Jest": "Testing & Quality",
    # Soft Skills
    "Leadership": "Soft Skills", "Communication": "Soft Skills", "Problem Solving": "Soft Skills",
    "Teamwork": "Soft Skills", "Critical Thinking": "Soft Skills", "Time Management": "Soft Skills",
}

ROLES = [
    "Backend Developer",
    "Frontend Developer",
    "Full Stack Developer",
    "Data Analyst",
    "Data Scientist",
    "DevOps Engineer",
    "QA Automation Engineer",
    "Software Engineer",
]

SENIORITY_LEVELS = ["Student/intern", "Junior", "Mid-level", "Senior"]

ROLE_CORE_SKILLS: dict[str, list[str]] = {
    "Backend Developer": ["Python", "FastAPI", "PostgreSQL", "Docker", "Redis", "Git", "Node.js", "Django", "SQL", "CI/CD"],
    "Frontend Developer": ["JavaScript", "TypeScript", "React", "HTML", "CSS", "Tailwind CSS", "Redux", "Git", "Next.js"],
    "Full Stack Developer": ["JavaScript", "TypeScript", "React", "Python", "FastAPI", "PostgreSQL", "Docker", "Git", "HTML", "CSS", "Redis"],
    "Data Analyst": ["Python", "SQL", "Excel", "Power BI", "Tableau", "Pandas", "NumPy", "Matplotlib", "Problem Solving"],
    "Data Scientist": ["Python", "Machine Learning", "Deep Learning", "PyTorch", "TensorFlow", "Scikit-learn", "Pandas", "NumPy", "SQL", "NLP"],
    "DevOps Engineer": ["AWS", "Docker", "Kubernetes", "Terraform", "Jenkins", "CI/CD", "Git", "Python", "Linux"],
    "QA Automation Engineer": ["Python", "Testing", "PyTest", "Selenium", "Postman", "Jest", "Git", "CI/CD", "JavaScript"],
    "Software Engineer": ["Java", "C++", "Python", "DSA", "Data Structures", "SQL", "Git", "Problem Solving", "Docker"],
}

SYNONYM_REPLACEMENTS: dict[str, str] = {
    "PostgreSQL": "Postgres",
    "Go": "Golang",
    "Kubernetes": "K8s",
    "React": "ReactJS",
    "Node.js": "Node js",
    "Vue.js": "Vue js",
    "Next.js": "Next js",
    "AWS": "AWS Cloud",
    "Scikit-learn": "sklearn",
    "Power BI": "PowerBI",
}

OUT_OF_DICTIONARY_SKILLS = [
    "Svelte", "Julia", "RustLang", "Solidity", "Elixir", "Haskell", "Prisma", "TailwindUI", "GraphQL", "WebAssembly"
]

CANONICAL_DEGREES = [
    "Bachelor of Technology",
    "B.Tech",
    "Bachelor of Engineering",
    "B.E",
    "B.Sc",
    "M.Sc",
    "M.Tech",
    "BCA",
    "MCA",
    "Diploma",
]

UNSUPPORTED_DEGREES = [
    "Bachelor of Science in Computing",
    "BS in Software Engineering",
    "Bachelor of Arts in Applied Math",
]

INSTITUTIONS = [
    "Apex Institute of Technology",
    "National Technical University",
    "Metropolitan Science Institute",
    "Horizon Engineering College",
    "Pinnacle Polytechnic School",
    "State Engineering University",
]

COMPANIES = [
    "Apex Global Software",
    "Nova Core Systems",
    "Vertex Cloud Solutions",
    "Synthetix Data Labs",
    "Blue Horizon Digital",
    "Quantum Dynamics Corp",
]

DISCLAIMER_HEADER = (
    "# =====================================================================\n"
    "# SYNTHETIC — NOT REAL RESEARCH DATA\n"
    "# Generated deterministically for controlled evaluation harness testing.\n"
    "# All personal details, names, and contact details are purely fictitious.\n"
    "# =====================================================================\n"
)


@dataclass
class ExperienceSpec:
    job_title: str
    company: str
    duration: str
    technologies: list[str]
    description: list[str]
    format_type: str = "same_line"  # "same_line", "next_line", "missing_date"


@dataclass
class ProjectSpec:
    title: str
    technologies: list[str]
    description: list[str]
    include_tech_line: bool = True


@dataclass
class EducationSpec:
    degree: Optional[str]
    institution: Optional[str]
    cgpa: Optional[str]
    duration: Optional[str]
    is_canonical_degree: bool = True


@dataclass
class ResumeSpec:
    resume_id: str
    candidate_name: str
    role: str
    seniority: str
    structure_type: str  # standard, alt_supported, unsupported, missing_sections, scrambled
    canonical_skills: list[str]
    surface_skills_in_skills_section: list[str]
    out_of_dict_skills: list[str]
    aspirational_skills: list[str]
    skills_in_projects_only: list[str]
    education: Optional[EducationSpec]
    experiences: list[ExperienceSpec]
    projects: list[ProjectSpec]
    certifications: list[str]


@dataclass
class JDSpec:
    jd_id: str
    title: str
    density_type: str  # high, medium, low, zero
    required_skills: list[str]
    preferred_skills: list[str]
    min_years_experience: int
    text_content: str


def group_skills_by_category(skills: list[str]) -> dict[str, list[str]]:
    grouped: dict[str, list[str]] = {}
    for s in skills:
        cat = SKILL_TO_CATEGORY.get(s, "Other")
        grouped.setdefault(cat, []).append(s)
    for cat in grouped:
        grouped[cat] = sorted(list(set(grouped[cat])))
    return grouped


# =====================================================================
# 1. 20 JOB DESCRIPTIONS SPECIFICATIONS
# =====================================================================
def create_synthetic_jds() -> list[JDSpec]:
    jds = []

    # 6 High-Density (10+ skills)
    jds.append(JDSpec(
        jd_id="syn_jd_01",
        title="Senior Fullstack Cloud Architect",
        density_type="high",
        required_skills=["Python", "FastAPI", "React", "TypeScript", "PostgreSQL", "Docker", "Kubernetes", "AWS"],
        preferred_skills=["Redis", "CI/CD", "Git", "Tailwind CSS"],
        min_years_experience=6,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Senior Fullstack Cloud Architect\n"
            "LOCATION: Remote / Global Operations\n"
            "EXPERIENCE REQUIRED: 6+ years\n\n"
            "ABOUT THE ROLE:\n"
            "We are seeking an experienced Senior Fullstack Cloud Architect to design and implement\n"
            "cloud-native web platforms. You will oversee end-to-end distributed system architecture.\n\n"
            "REQUIRED TECHNICAL SKILLS:\n"
            "- Core Languages & Frameworks: Python, FastAPI, React, TypeScript\n"
            "- Data Storage: PostgreSQL, Redis\n"
            "- Cloud & Containers: AWS, Docker, Kubernetes\n\n"
            "PREFERRED SKILLS:\n"
            "- CI/CD, Git version control, Tailwind CSS\n\n"
            "RESPONSIBILITIES:\n"
            "- Architect highly scalable asynchronous microservices.\n"
            "- Lead development of enterprise single-page applications.\n"
            "- Manage containerized production Kubernetes clusters on AWS.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_02",
        title="Lead Data Scientist",
        density_type="high",
        required_skills=["Python", "Machine Learning", "Deep Learning", "PyTorch", "TensorFlow", "Scikit-learn", "Pandas", "NumPy", "SQL"],
        preferred_skills=["NLP", "Docker", "Git"],
        min_years_experience=6,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Lead Data Scientist\n"
            "LOCATION: Synthetic Tech Hub\n"
            "EXPERIENCE REQUIRED: 6+ years\n\n"
            "JOB DESCRIPTION:\n"
            "Join our predictive modeling division to lead development of advanced machine learning pipelines.\n\n"
            "REQUIRED SKILLS:\n"
            "- Machine Learning, Deep Learning, PyTorch, TensorFlow, Scikit-learn\n"
            "- Python, Pandas, NumPy, SQL\n\n"
            "PREFERRED QUALIFICATIONS:\n"
            "- Experience with NLP, containerization with Docker, and Git workflows.\n\n"
            "DUTIES:\n"
            "- Train, tune, and evaluate deep neural networks.\n"
            "- Transform raw business tables using SQL into feature-engineered datasets.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_03",
        title="Principal Backend Distributed Systems Engineer",
        density_type="high",
        required_skills=["Go", "Python", "Java", "Docker", "Kubernetes", "PostgreSQL", "Redis", "AWS"],
        preferred_skills=["MongoDB", "CI/CD", "Git", "Data Structures"],
        min_years_experience=7,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Principal Backend Distributed Systems Engineer\n"
            "LOCATION: Synthetic Hub\n"
            "EXPERIENCE REQUIRED: 7+ years\n\n"
            "TECHNICAL REQUIREMENTS:\n"
            "- Strong command of Go, Python, and Java for concurrent service development.\n"
            "- Database expertise: PostgreSQL, Redis, MongoDB.\n"
            "- Infrastructure: Docker, Kubernetes, AWS, CI/CD, Git.\n"
            "- Solid mastery of Data Structures and system design.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_04",
        title="Senior Cloud DevOps & Site Reliability Engineer",
        density_type="high",
        required_skills=["AWS", "Azure", "Docker", "Kubernetes", "Terraform", "Jenkins", "CI/CD", "Python"],
        preferred_skills=["Git", "Linux"],
        min_years_experience=5,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Senior Cloud DevOps & Site Reliability Engineer\n"
            "LOCATION: Remote\n"
            "EXPERIENCE: 5+ years\n\n"
            "REQUIREMENTS:\n"
            "- Cloud platforms: AWS, Azure\n"
            "- Infrastructure as Code: Terraform, Docker, Kubernetes\n"
            "- Automation & Scripting: Jenkins, CI/CD, Python, Git\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_05",
        title="Senior Data Platform Engineer",
        density_type="high",
        required_skills=["Python", "SQL", "PostgreSQL", "MySQL", "MongoDB", "Redis", "Docker", "AWS"],
        preferred_skills=["Kubernetes", "Git", "CI/CD"],
        min_years_experience=5,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Senior Data Platform Engineer\n"
            "LOCATION: Hybrid\n"
            "REQUIREMENTS:\n"
            "- Languages & Storage: Python, SQL, PostgreSQL, MySQL, MongoDB, Redis\n"
            "- Cloud Infrastructure: Docker, AWS, Kubernetes, Git, CI/CD\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_06",
        title="Full Stack Web & API Specialist",
        density_type="high",
        required_skills=["JavaScript", "TypeScript", "React", "Node.js", "Express.js", "HTML", "CSS", "Tailwind CSS", "MongoDB"],
        preferred_skills=["Git", "Postman"],
        min_years_experience=4,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Full Stack Web & API Specialist\n"
            "LOCATION: On-site\n"
            "REQUIREMENTS:\n"
            "- Frontend: JavaScript, TypeScript, React, HTML, CSS, Tailwind CSS\n"
            "- Backend: Node.js, Express.js, MongoDB\n"
            "- Testing & Versioning: Git, Postman\n"
        )
    ))

    # 8 Medium-Density (5-9 skills)
    jds.append(JDSpec(
        jd_id="syn_jd_07",
        title="Mid-Level Python Backend Developer",
        density_type="medium",
        required_skills=["Python", "Django", "PostgreSQL", "Docker", "Git"],
        preferred_skills=["Redis", "CI/CD"],
        min_years_experience=3,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Mid-Level Python Backend Developer\n"
            "REQUIREMENTS:\n"
            "- Proficient with Python and Django web framework.\n"
            "- Relational databases: PostgreSQL.\n"
            "- Tooling: Docker, Git, Redis, CI/CD.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_08",
        title="Frontend React Engineer",
        density_type="medium",
        required_skills=["JavaScript", "TypeScript", "React", "HTML", "CSS"],
        preferred_skills=["Redux", "Tailwind CSS"],
        min_years_experience=2,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Frontend React Engineer\n"
            "REQUIREMENTS:\n"
            "- Core stack: JavaScript, TypeScript, React, HTML, CSS.\n"
            "- State management: Redux. Styling: Tailwind CSS.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_09",
        title="Business Data Analyst",
        density_type="medium",
        required_skills=["Python", "SQL", "Excel", "Power BI", "Tableau"],
        preferred_skills=["Problem Solving"],
        min_years_experience=2,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Business Data Analyst\n"
            "REQUIREMENTS:\n"
            "- Hands-on analytics: Python, SQL, Excel.\n"
            "- Dashboard creation: Power BI, Tableau.\n"
            "- Soft skills: Problem Solving.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_10",
        title="QA Automation Engineer",
        density_type="medium",
        required_skills=["Python", "Testing", "Selenium", "PyTest", "Postman"],
        preferred_skills=["Git", "CI/CD"],
        min_years_experience=3,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: QA Automation Engineer\n"
            "REQUIREMENTS:\n"
            "- Quality engineering: Testing, Selenium, PyTest, Postman.\n"
            "- Scripting: Python, Git, CI/CD.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_11",
        title="Junior Software Engineer",
        density_type="medium",
        required_skills=["Java", "SQL", "Git", "DSA", "Problem Solving"],
        preferred_skills=["C++"],
        min_years_experience=1,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Junior Software Engineer\n"
            "REQUIREMENTS:\n"
            "- Fundamental programming: Java, SQL, Git, DSA, Problem Solving, C++.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_12",
        title="Cloud Infrastructure Associate",
        density_type="medium",
        required_skills=["AWS", "Docker", "Terraform", "CI/CD", "Git"],
        preferred_skills=["Python"],
        min_years_experience=2,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Cloud Infrastructure Associate\n"
            "REQUIREMENTS:\n"
            "- Cloud deployment: AWS, Docker, Terraform, CI/CD, Git, Python.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_13",
        title="Machine Learning Engineer",
        density_type="medium",
        required_skills=["Python", "PyTorch", "Scikit-learn", "Docker", "Git"],
        preferred_skills=["FastAPI"],
        min_years_experience=3,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Machine Learning Engineer\n"
            "REQUIREMENTS:\n"
            "- Model building: Python, PyTorch, Scikit-learn.\n"
            "- Serving & Deploy: Docker, Git, FastAPI.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_14",
        title="Web Application Developer",
        density_type="medium",
        required_skills=["PHP", "Laravel", "MySQL", "HTML", "CSS", "JavaScript"],
        preferred_skills=["Git"],
        min_years_experience=2,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Web Application Developer\n"
            "REQUIREMENTS:\n"
            "- Stack: PHP, Laravel, MySQL, HTML, CSS, JavaScript, Git.\n"
        )
    ))

    # 4 Low-Density (1-4 skills)
    jds.append(JDSpec(
        jd_id="syn_jd_15",
        title="Entry Level IT Support Coordinator",
        density_type="low",
        required_skills=["Excel", "Problem Solving", "Communication"],
        preferred_skills=[],
        min_years_experience=0,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Entry Level IT Support Coordinator\n"
            "REQUIREMENTS:\n"
            "- Knowledge of Excel spreadsheet tracking.\n"
            "- Strong Problem Solving and Communication skills.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_16",
        title="Junior Python Scripting Assistant",
        density_type="low",
        required_skills=["Python", "Git"],
        preferred_skills=[],
        min_years_experience=1,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Junior Python Scripting Assistant\n"
            "REQUIREMENTS:\n"
            "- Basic automation with Python and version tracking with Git.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_17",
        title="Junior SQL Reporting Associate",
        density_type="low",
        required_skills=["SQL", "MySQL"],
        preferred_skills=[],
        min_years_experience=1,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Junior SQL Reporting Associate\n"
            "REQUIREMENTS:\n"
            "- Writing basic queries in SQL and managing tables in MySQL.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_18",
        title="Static Web Content Assistant",
        density_type="low",
        required_skills=["HTML", "CSS"],
        preferred_skills=[],
        min_years_experience=0,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Static Web Content Assistant\n"
            "REQUIREMENTS:\n"
            "- Formatting static articles using semantic HTML and CSS styling.\n"
        )
    ))

    # 2 Zero Recognized Dictionary Skills
    jds.append(JDSpec(
        jd_id="syn_jd_19",
        title="Executive Facilities & Procurement Director",
        density_type="zero",
        required_skills=[],
        preferred_skills=[],
        min_years_experience=8,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Executive Facilities & Procurement Director\n"
            "LOCATION: Synthetic Corporate Headquarters\n"
            "EXPERIENCE: 8+ years\n\n"
            "ROLE OVERVIEW:\n"
            "Direct corporate real estate negotiations, commercial lease agreements,\n"
            "physical facility maintenance contracts, and third-party supplier logistics.\n"
            "Must maintain strict municipal health code compliance and oversee building security.\n"
        )
    ))

    jds.append(JDSpec(
        jd_id="syn_jd_20",
        title="Strategic Public Relations & Media Diplomat",
        density_type="zero",
        required_skills=[],
        preferred_skills=[],
        min_years_experience=7,
        text_content=(
            f"{DISCLAIMER_HEADER}\n"
            "JOB TITLE: Strategic Public Relations & Media Diplomat\n"
            "LOCATION: Synthetic Media Bureau\n"
            "EXPERIENCE: 7+ years\n\n"
            "ROLE OVERVIEW:\n"
            "Spearhead high-visibility press briefings, crisis mitigation messaging,\n"
            "investor conference announcements, and editorial interviews.\n"
            "Requires diplomatic charisma, television presence, and executive stakeholder briefing skills.\n"
        )
    ))

    return jds


# =====================================================================
# 2. 100 RESUMES SPECIFICATION BUILDER
# =====================================================================
def create_synthetic_resume_specs(rng: random.Random) -> list[ResumeSpec]:
    specs: list[ResumeSpec] = []

    structures = [
        "standard",
        "alt_supported",
        "unsupported",
        "missing_sections",
        "scrambled",
    ]

    for index in range(1, 101):
        resume_id = f"syn_resume_{index:03d}"
        candidate_name = f"SyntheticCandidate_{index:03d}"
        role = ROLES[(index - 1) % len(ROLES)]
        seniority = SENIORITY_LEVELS[(index - 1) % len(SENIORITY_LEVELS)]
        structure_type = structures[(index - 1) % len(structures)]

        core_pool = ROLE_CORE_SKILLS[role]
        num_skills = {
            "Student/intern": rng.randint(4, 6),
            "Junior": rng.randint(5, 7),
            "Mid-level": rng.randint(6, 9),
            "Senior": rng.randint(8, 11),
        }[seniority]

        selected_skills = rng.sample(core_pool, min(num_skills, len(core_pool)))
        # occasionally add adjacent skills
        if rng.random() > 0.4:
            other_skills = [s for s in SKILL_TO_CATEGORY.keys() if s not in selected_skills]
            selected_skills.append(rng.choice(other_skills))
        selected_skills = sorted(list(set(selected_skills)))

        # Introduce surface form variations for testing
        surface_skills = []
        for s in selected_skills:
            if s in SYNONYM_REPLACEMENTS and rng.random() > 0.5:
                surface_skills.append(SYNONYM_REPLACEMENTS[s])
            elif rng.random() < 0.15:
                surface_skills.append(s.lower())
            elif rng.random() < 0.05:
                surface_skills.append(s.upper())
            else:
                surface_skills.append(s)

        out_of_dict = []
        if rng.random() > 0.6:
            out_of_dict.append(rng.choice(OUT_OF_DICTIONARY_SKILLS))

        aspirational = []
        if rng.random() > 0.7:
            unused = [s for s in SKILL_TO_CATEGORY if s not in selected_skills]
            aspirational.append(rng.choice(unused))

        skills_in_proj_only = []
        if rng.random() > 0.6 and len(selected_skills) > 4:
            skills_in_proj_only.append(selected_skills.pop())

        # Build Education
        if structure_type == "missing_sections" and rng.random() > 0.5:
            edu = None
        else:
            is_canonical = rng.random() > 0.2
            degree = rng.choice(CANONICAL_DEGREES) if is_canonical else rng.choice(UNSUPPORTED_DEGREES)
            inst = rng.choice(INSTITUTIONS)
            cgpa = f"{rng.uniform(3.2, 4.0):.2f}" if rng.random() > 0.3 else None
            grad_year = rng.randint(2015, 2024)
            duration = f"{grad_year - 4} - {grad_year}" if rng.random() > 0.2 else f"{grad_year}"
            edu = EducationSpec(
                degree=degree,
                institution=inst,
                cgpa=cgpa,
                duration=duration,
                is_canonical_degree=is_canonical,
            )

        # Build Experiences
        experiences: list[ExperienceSpec] = []
        if structure_type == "missing_sections" and seniority == "Student/intern":
            pass  # no experience
        else:
            num_jobs = {"Student/intern": 1, "Junior": 1, "Mid-level": 2, "Senior": 3}[seniority]
            start_year = 2024 - ({"Student/intern": 1, "Junior": 2, "Mid-level": 5, "Senior": 9}[seniority])
            for j in range(num_jobs):
                dur_start = f"Jan {start_year + j * 2}"
                dur_end = "Present" if j == num_jobs - 1 else f"Dec {start_year + j * 2 + 1}"
                fmt = "next_line" if (rng.random() < 0.15) else "same_line"
                tech_used = rng.sample(selected_skills, min(len(selected_skills), rng.randint(2, 4)))
                desc = [
                    f"Architected core {role} features utilizing {tech_used[0]}.",
                    f"Enhanced system efficiency and automated processing using {tech_used[-1]}.",
                ]
                experiences.append(ExperienceSpec(
                    job_title=f"{'Senior ' if seniority == 'Senior' and j == num_jobs - 1 else ''}{role}",
                    company=rng.choice(COMPANIES),
                    duration=f"{dur_start} - {dur_end}",
                    technologies=tech_used,
                    description=desc,
                    format_type=fmt,
                ))

        # Build Projects
        projects: list[ProjectSpec] = []
        if not (structure_type == "missing_sections" and rng.random() > 0.6):
            num_proj = rng.randint(1, 2)
            for p in range(num_proj):
                p_tech = rng.sample(selected_skills, min(len(selected_skills), rng.randint(2, 3)))
                if skills_in_proj_only:
                    p_tech.extend(skills_in_proj_only)
                p_tech = sorted(list(set(p_tech)))
                projects.append(ProjectSpec(
                    title=f"Synthetic Platform Project {p + 1}",
                    technologies=p_tech,
                    description=[
                        f"Constructed scalable service processing data with {p_tech[0]}.",
                        f"Integrated automated pipelines and testing with {p_tech[-1]}.",
                    ],
                    include_tech_line=rng.random() > 0.3,
                ))

        # Build Certifications
        certs = []
        if seniority in ("Mid-level", "Senior") and structure_type != "missing_sections":
            certs = [
                f"{rng.choice(['AWS', 'Docker', 'PostgreSQL'])} Certified Professional",
                "Certified Systems Developer",
            ][:rng.randint(1, 2)]

        specs.append(ResumeSpec(
            resume_id=resume_id,
            candidate_name=candidate_name,
            role=role,
            seniority=seniority,
            structure_type=structure_type,
            canonical_skills=selected_skills,
            surface_skills_in_skills_section=surface_skills,
            out_of_dict_skills=out_of_dict,
            aspirational_skills=aspirational,
            skills_in_projects_only=skills_in_proj_only,
            education=edu,
            experiences=experiences,
            projects=projects,
            certifications=certs,
        ))

    return specs


# =====================================================================
# 3. TEXT RENDERING (TXT, PDF, DOCX)
# =====================================================================
def render_resume_text(spec: ResumeSpec) -> str:
    lines = [
        DISCLAIMER_HEADER,
        spec.candidate_name.upper(),
        f"Role: {spec.role} | Experience: {spec.seniority}",
        "synthetic.candidate@example.org | (555) 019-8234 | Synthetic City, SC",
        "",
    ]

    headings = {
        "summary": "SUMMARY",
        "skills": "SKILLS",
        "experience": "EXPERIENCE",
        "education": "EDUCATION",
        "projects": "PROJECTS",
        "certifications": "CERTIFICATIONS",
    }
    if spec.structure_type == "alt_supported":
        headings = {
            "summary": "PROFESSIONAL SUMMARY",
            "skills": "TECHNICAL EXPERTISE",
            "experience": "WORK EXPERIENCE",
            "education": "ACADEMIC BACKGROUND",
            "projects": "KEY PROJECTS",
            "certifications": "COURSES",
        }
    elif spec.structure_type == "unsupported":
        headings = {
            "summary": "MY PHILOSOPHY",
            "skills": "CORE COMPETENCIES & TOOLKIT",
            "experience": "CAREER TRAJECTORY",
            "education": "SCHOOLING AND DEGREES",
            "projects": "CREATIVE ENDEAVORS",
            "certifications": "BADGES AND PAPERS",
        }

    sections_content: dict[str, list[str]] = {}

    # Summary
    summary_text = f"Dedicated {spec.seniority} {spec.role} with proven ability to build scalable software."
    if spec.aspirational_skills:
        summary_text += f" Currently exploring and learning {', '.join(spec.aspirational_skills)}."
    sections_content["summary"] = [headings["summary"], "-" * 30, summary_text, ""]

    # Skills
    displayed_skills = list(spec.surface_skills_in_skills_section)
    if spec.out_of_dict_skills:
        displayed_skills.extend(spec.out_of_dict_skills)
    sections_content["skills"] = [
        headings["skills"],
        "-" * 30,
        f"Technical Skills: {', '.join(displayed_skills)}",
        "",
    ]

    # Experience
    exp_lines = [headings["experience"], "-" * 30]
    if spec.experiences:
        for exp in spec.experiences:
            if exp.format_type == "same_line":
                exp_lines.append(f"{exp.job_title}   {exp.duration}")
                exp_lines.append(exp.company)
            else:
                # Next line format (triggers known extractor failure)
                exp_lines.append(exp.job_title)
                exp_lines.append(exp.company)
                exp_lines.append(exp.duration)
            for bullet in exp.description:
                exp_lines.append(f"• {bullet}")
            exp_lines.append("")
    else:
        exp_lines.append("No previous formal industry employment.\n")
    sections_content["experience"] = exp_lines

    # Education
    edu_lines = [headings["education"], "-" * 30]
    if spec.education and spec.education.degree:
        edu_lines.append(spec.education.degree)
        edu_lines.append(spec.education.institution or "Synthetic Institute")
        if spec.education.duration:
            edu_lines.append(spec.education.duration)
        if spec.education.cgpa:
            edu_lines.append(f"CGPA: {spec.education.cgpa}")
        edu_lines.append("")
    else:
        edu_lines.append("Self-taught practical background.\n")
    sections_content["education"] = edu_lines

    # Projects
    proj_lines = [headings["projects"], "-" * 30]
    if spec.projects:
        for p in spec.projects:
            proj_lines.append(p.title)
            if p.include_tech_line and p.technologies:
                proj_lines.append(f"Technologies: {', '.join(p.technologies)}")
            for b in p.description:
                proj_lines.append(f"• {b}")
            proj_lines.append("")
    else:
        proj_lines.append("Portfolio available upon request.\n")
    sections_content["projects"] = proj_lines

    # Certifications
    cert_lines = [headings["certifications"], "-" * 30]
    if spec.certifications:
        for c in spec.certifications:
            cert_lines.append(c)
        cert_lines.append("")
    else:
        cert_lines.append("None.\n")
    sections_content["certifications"] = cert_lines

    # Ordering
    order = ["summary", "skills", "experience", "education", "projects", "certifications"]
    if spec.structure_type == "scrambled":
        order = ["education", "projects", "skills", "experience", "certifications", "summary"]

    for sec in order:
        lines.extend(sections_content[sec])

    return "\n".join(lines)


def render_resume_pdf(spec: ResumeSpec, text_content: str, output_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)  # A4 size
    rect = pymupdf.Rect(40, 40, 555, 802)
    # PyMuPDF textbox handles newlines and basic layout
    page.insert_textbox(rect, text_content, fontsize=9.5, fontname="helv", lineheight=1.2)
    doc.save(str(output_path))
    doc.close()


def render_resume_docx(spec: ResumeSpec, text_content: str, output_path: Path) -> None:
    doc = docx.Document()
    for line in text_content.split("\n"):
        line_clean = line.strip()
        if not line_clean:
            continue
        if line_clean.startswith("# =="):
            continue
        if line_clean.isupper() and len(line_clean) < 35:
            doc.add_heading(line_clean, level=2)
        elif line_clean.startswith("•"):
            doc.add_paragraph(line_clean.lstrip("• "), style="List Bullet")
        else:
            doc.add_paragraph(line_clean)
    doc.save(str(output_path))


# =====================================================================
# 4. GOLD ANNOTATIONS (DERIVED STRICTLY FROM SPECS BEFORE PARSING)
# =====================================================================
def compute_gold_resume_annotation(spec: ResumeSpec) -> dict[str, Any]:
    # True skills: All canonical skills intentionally put into the resume
    all_true_skills = set(spec.canonical_skills)
    all_true_skills.update(spec.skills_in_projects_only)

    gold_skills_grouped = group_skills_by_category(list(all_true_skills))

    gold_edu = {
        "degree": spec.education.degree if (spec.education and spec.education.degree) else None,
        "institution": spec.education.institution if (spec.education and spec.education.institution) else None,
        "cgpa": spec.education.cgpa if (spec.education and spec.education.cgpa) else None,
        "duration": spec.education.duration if (spec.education and spec.education.duration) else None,
    }

    gold_exp = []
    for exp in spec.experiences:
        gold_exp.append({
            "job_title": exp.job_title,
            "company": exp.company,
            "duration": exp.duration,
            "technologies": sorted(list(set(exp.technologies))),
            "description": exp.description,
        })

    gold_proj = []
    for p in spec.projects:
        gold_proj.append({
            "title": p.title,
            "technologies": sorted(list(set(p.technologies))),
            "description": p.description,
        })

    return {
        "resume_id": spec.resume_id,
        "file_path": f"resumes/{spec.resume_id}.txt",
        "skills": gold_skills_grouped,
        "education": gold_edu,
        "experience": gold_exp,
        "projects": gold_proj,
        "certifications": spec.certifications,
        "notes": f"SYNTHETIC — NOT REAL RESEARCH DATA. Spec: role={spec.role}, seniority={spec.seniority}, structure={spec.structure_type}.",
    }


def compute_gold_match(
    spec: ResumeSpec, jd: JDSpec
) -> tuple[dict[str, Any], float]:
    candidate_skills = set(spec.canonical_skills) | set(spec.skills_in_projects_only)
    jd_skills = set(jd.required_skills) | set(jd.preferred_skills)

    matched = sorted(list(candidate_skills & jd_skills))
    missing = sorted(list(jd_skills - candidate_skills))
    relevant = sorted(list(jd_skills))

    # Compute Expert Relevance Score deterministically (0-100 rubric)
    if jd.density_type == "zero":
        # Candidate evaluated on general qualification and role affinity
        if spec.seniority == "Senior":
            expert_score = 45.0
        elif spec.seniority == "Mid-level":
            expert_score = 35.0
        else:
            expert_score = 20.0
    else:
        coverage = len(matched) / len(jd_skills) if jd_skills else 1.0
        skill_component = coverage * 50.0  # up to 50 pts

        # Seniority alignment (up to 25 pts)
        seniority_years = {"Student/intern": 1, "Junior": 2, "Mid-level": 4, "Senior": 8}[spec.seniority]
        if seniority_years >= jd.min_years_experience:
            sen_component = 25.0
        else:
            gap = jd.min_years_experience - seniority_years
            sen_component = max(5.0, 25.0 - (gap * 6.0))

        # Role alignment (up to 15 pts)
        role_component = 15.0 if jd.title.lower().startswith(spec.role.lower()[:7]) else 5.0

        # Education & certifications bonus (up to 10 pts)
        edu_component = 5.0 if (spec.education and spec.education.degree) else 0.0
        cert_component = min(5.0, len(spec.certifications) * 2.5)

        expert_score = round(min(100.0, skill_component + sen_component + role_component + edu_component + cert_component), 1)

    match_record = {
        "match_id": f"{spec.resume_id}__{jd.jd_id}",
        "resume_id": spec.resume_id,
        "jd_id": jd.jd_id,
        "gold_relevant_skills": relevant,
        "gold_matched_skills": matched,
        "gold_missing_skills": missing,
        "notes": f"SYNTHETIC — NOT REAL RESEARCH DATA. Match {spec.role} vs {jd.title}.",
    }

    return match_record, expert_score


# =====================================================================
# 5. ORCHESTRATOR
# =====================================================================
def generate_dataset(output_dir: Path, seed: int = 42) -> None:
    rng = random.Random(seed)

    resumes_dir = output_dir / "resumes"
    jds_dir = output_dir / "job_descriptions"
    annotations_dir = output_dir / "annotations"
    scores_dir = output_dir / "expert_scores"
    rankings_dir = output_dir / "expert_rankings"

    for d in (resumes_dir, jds_dir, annotations_dir, scores_dir, rankings_dir):
        d.mkdir(parents=True, exist_ok=True)

    print("1. Generating 20 Synthetic Job Descriptions...")
    jds = create_synthetic_jds()
    for jd in jds:
        jd_file = jds_dir / f"{jd.jd_id}.txt"
        jd_file.write_text(jd.text_content, encoding="utf-8")

    print("2. Generating 100 Synthetic Resumes (TXT, and first 25 as PDF/DOCX)...")
    specs = create_synthetic_resume_specs(rng)
    resume_annotations: list[dict] = []

    for idx, spec in enumerate(specs):
        text = render_resume_text(spec)
        txt_path = resumes_dir / f"{spec.resume_id}.txt"
        txt_path.write_text(text, encoding="utf-8")

        # Multi-format rendering for first 25
        if idx < 25:
            pdf_path = resumes_dir / f"{spec.resume_id}.pdf"
            docx_path = resumes_dir / f"{spec.resume_id}.docx"
            render_resume_pdf(spec, text, pdf_path)
            render_resume_docx(spec, text, docx_path)

        resume_annotations.append(compute_gold_resume_annotation(spec))

    # Save resume annotations JSONL
    (annotations_dir / "resume_annotations.jsonl").write_text(
        "\n".join(json.dumps(r) for r in resume_annotations) + "\n",
        encoding="utf-8",
    )

    print("3. Generating 320 Resume-JD Pairs & Gold Match/Score Annotations...")
    match_annotations: list[dict] = []
    expert_score_records: list[dict] = []
    expert_score_csv_rows = [
        ["# SYNTHETIC — NOT REAL RESEARCH DATA"],
        ["resume_id", "jd_id", "expert_relevance_score", "evaluator_id", "scoring_rationale"],
    ]

    cohort_rankings: list[dict] = []
    ranking_csv_rows = [
        ["# SYNTHETIC — NOT REAL RESEARCH DATA"],
        ["group_id", "jd_id", "candidate_id", "expert_rank", "evaluator_id"],
    ]

    # Assign 16 resumes to each JD across match tiers
    for jd in jds:
        # Sort specs by heuristic fit to pick high, partial, and low match candidates
        def fit_score(s: ResumeSpec) -> float:
            c_skills = set(s.canonical_skills)
            jd_s = set(jd.required_skills) | set(jd.preferred_skills)
            return len(c_skills & jd_s)

        sorted_specs = sorted(specs, key=fit_score, reverse=True)
        # 5 high match, 5 partial match, 4 low match, 2 zero match
        high_group = sorted_specs[:5]
        partial_group = sorted_specs[25:30]
        low_group = sorted_specs[55:59]
        zero_group = sorted_specs[-2:]
        cohort = high_group + partial_group + low_group + zero_group  # total = 16

        cohort_scores: list[tuple[str, float]] = []

        for spec in cohort:
            m_rec, exp_score = compute_gold_match(spec, jd)
            match_annotations.append(m_rec)

            score_rec = {
                "resume_id": spec.resume_id,
                "jd_id": jd.jd_id,
                "expert_relevance_score": exp_score,
                "evaluator_id": "synthetic_ground_truth_oracle",
                "scoring_rationale": f"SYNTHETIC — NOT REAL RESEARCH DATA. Formulaic rubric match: {exp_score}/100.",
            }
            expert_score_records.append(score_rec)
            expert_score_csv_rows.append([
                spec.resume_id,
                jd.jd_id,
                str(exp_score),
                "synthetic_ground_truth_oracle",
                f"Synthetic score {exp_score}",
            ])
            cohort_scores.append((spec.resume_id, exp_score))

        # Rank cohort: highest score = rank 1
        cohort_scores.sort(key=lambda item: (-item[1], item[0]))
        rank_dict: dict[str, int] = {}
        for rank_num, (cand_id, _) in enumerate(cohort_scores, start=1):
            rank_dict[cand_id] = rank_num
            ranking_csv_rows.append([
                f"cohort_{jd.jd_id}",
                jd.jd_id,
                cand_id,
                str(rank_num),
                "synthetic_ground_truth_oracle",
            ])

        cohort_rankings.append({
            "group_id": f"cohort_{jd.jd_id}",
            "jd_id": jd.jd_id,
            "candidate_ids": [c[0] for c in cohort_scores],
            "expert_rank": rank_dict,
            "evaluator_id": "synthetic_ground_truth_oracle",
            "notes": f"SYNTHETIC — NOT REAL RESEARCH DATA. Cohort for {jd.title}.",
        })

    # Save match annotations JSONL
    (annotations_dir / "match_annotations.jsonl").write_text(
        "\n".join(json.dumps(r) for r in match_annotations) + "\n",
        encoding="utf-8",
    )

    # Save expert scores JSONL and CSV
    (scores_dir / "expert_scores.jsonl").write_text(
        "\n".join(json.dumps(r) for r in expert_score_records) + "\n",
        encoding="utf-8",
    )
    with (scores_dir / "expert_scores.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(expert_score_csv_rows)

    # Save expert rankings JSONL and CSV
    (rankings_dir / "expert_rankings.jsonl").write_text(
        "\n".join(json.dumps(r) for r in cohort_rankings) + "\n",
        encoding="utf-8",
    )
    with (rankings_dir / "expert_rankings.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(ranking_csv_rows)

    print(f"\n[SUCCESS] Generated Synthetic Dataset in {output_dir}:")
    print(f"  - Resumes: {len(specs)} (.txt), 25 (.pdf), 25 (.docx)")
    print(f"  - Job Descriptions: {len(jds)}")
    print(f"  - Resume-JD Evaluated Pairs: {len(match_annotations)}")
    print(f"  - Gold Resume Annotations: {len(resume_annotations)}")
    print(f"  - Cohort Ranking Groups: {len(cohort_rankings)}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        default=str(Path(__file__).resolve().parent / "data"),
        help="Target directory for synthetic data output",
    )
    parser.add_argument("--seed", type=int, default=42, help="Random seed for deterministic generation")
    args = parser.parse_args()

    out = Path(args.output_dir)
    generate_dataset(out, seed=args.seed)


if __name__ == "__main__":
    main()
